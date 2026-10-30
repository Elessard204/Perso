"""Crée un jeu de données fictif pour essayer l'outil sans toucher aux vrais clients.

Produit demo/clients_demo.xlsx et demo/exports/2026-10/*.pdf, avec quelques
problèmes volontaires (rapport manquant, fichier mal nommé, courriel invalide)
pour montrer ce que l'outil bloque.
"""

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

ICI = Path(__file__).resolve().parent
MOIS = "2026-10"
DOCS = ["etat_resultats", "bilan", "sommaire", "grand_livre"]
TITRES = {"etat_resultats": "Etat des resultats", "bilan": "Bilan",
          "sommaire": "Sommaire mensuel des transactions", "grand_livre": "Grand livre"}

CLIENTS = [
    # code, nom, contact, courriels clients, gestionnaire, courriel gestionnaires, courriel direction, langue, actif
    ("C001", "Gestion Boréale inc.", "Julie", "julie.gagnon@exemple.com", "Marc Bédard", "mbedard@progim.com", "", "FR", "Oui"),
    ("C002", "Saint-Laurent Holdings", "David", "david.cohen@exemple.com; comptable@exemple.com", "Sophie Roy", "sroy@progim.com", "elessard@progim.com", "EN", "Oui"),
    ("C003", "Immeubles Rivière-Noire", "Pierre", "pierre.lavoie@exemple.com", "Marc Bédard", "mbedard@progim.com", "", "FR", "Oui"),
    ("C004", "Famille Tremblay", "Nathalie", "n.tremblay@exemple.com", "Sophie Roy", "sroy@progim.com", "", "FR", "Oui"),
    ("C005", "Société Champlain", "Louis", "compta@champlain", "Marc Bédard", "mbedard@progim.com", "", "FR", "Oui"),
    ("C006", "Ancien client ltée", "Paul", "paul@exemple.com", "Sophie Roy", "sroy@progim.com", "", "FR", "Non"),
    ("C007", "Groupe Laurentides", "Isabelle", "isabelle.cote@exemple.com", "Marc Bédard", "mbedard@progim.com", "", "FR", "Oui"),
]

ENTITES = [
    # # client Hopem, nom, client, début d'exercice (mois), gestionnaire propre à l'entité, documents requis
    ("1001", "Boréale Immeubles inc.", "C001", 1, None, DOCS),
    ("1002", "9123-4567 Québec inc.", "C001", 11, ("Karine Ouellet", "kouellet@progim.com"), DOCS),
    ("2001", "SLH Residential LP", "C002", 4, None, ["etat_resultats", "bilan"]),
    ("3001", "Rivière-Noire Habitation inc.", "C003", 1, None, DOCS),
    ("4001", "Tremblay Immobilier inc.", "C004", 1, None, DOCS),
    ("4002", "Résidences Tremblay S.E.C.", "C004", 1, None, ["etat_resultats", "bilan", "sommaire"]),
    ("5001", "Champlain Logements inc.", "C005", 5, None, DOCS),
    ("6001", "Ancien client ltée", "C006", 1, None, DOCS),
    ("7001", "Laurentides Multilogements inc.", "C007", 11, None, ["etat_resultats", "bilan", "grand_livre"]),
]


def creer_excel() -> None:
    classeur = Workbook()
    entete = Font(bold=True, color="FFFFFF")
    fond = PatternFill("solid", fgColor="1F5C4D")

    onglets = [
        ("Clients", ["code_client", "nom_client", "nom_contact", "courriel_clients", "gestionnaire_nom",
                     "courriel_gestionnaires", "courriel_direction", "langue", "actif"], CLIENTS),
        ("Entites", ["code_entite", "nom_entite", "code_client", "debut_exercice", "gestionnaire_nom",
                     "courriel_gestionnaire", *DOCS, "actif"],
         [(c, n, cl, debut, *(g or ("", "")), *["Oui" if d in docs else "Non" for d in DOCS], "Oui")
          for c, n, cl, debut, g, docs in ENTITES]),
    ]
    for i, (nom, colonnes, lignes) in enumerate(onglets):
        feuille = classeur.active if i == 0 else classeur.create_sheet()
        feuille.title = nom
        feuille.append(colonnes)
        for ligne in lignes:
            feuille.append(list(ligne))
        for cellule in feuille[1]:
            cellule.font, cellule.fill = entete, fond
        for colonne in feuille.columns:
            feuille.column_dimensions[colonne[0].column_letter].width = max(len(str(c.value or "")) for c in colonne) + 4
        table = Table(displayName=f"T_{nom}", ref=feuille.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        feuille.add_table(table)
        feuille.freeze_panes = "A2"
    classeur.save(ICI / "clients_demo.xlsx")


def creer_pdf(chemin: Path, code: str, nom: str, doc: str) -> None:
    pdf = canvas.Canvas(str(chemin), pagesize=letter)
    pages = 3 if doc == "grand_livre" else 1
    for page in range(pages):
        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(50, 740, f"{TITRES[doc]} - {MOIS}")
        pdf.setFont("Helvetica", 10)
        pdf.drawString(50, 722, f"No client Hopem {code} - {nom}  (donnees fictives)   page {page + 1}/{pages}")
        y = 690
        for _ in range(25):
            pdf.drawString(50, y, f"{random.randint(1000, 5999)}  Poste fictif")
            pdf.drawRightString(560, y, f"{random.uniform(-20000, 90000):,.2f} $")
            y -= 22
        pdf.showPage()
    pdf.save()


def creer_exports() -> None:
    dossier = ICI / "exports" / MOIS
    dossier.mkdir(parents=True, exist_ok=True)
    for ancien in dossier.glob("*.pdf"):
        ancien.unlink()
    random.seed(2026)
    for code, nom, client, debut, _, docs in ENTITES:
        for doc in docs:
            if code == "3001" and doc == "bilan":
                continue  # oubli d'export : le client C003 sera bloqué
            mois = "Septembre" if (code, doc) == ("7001", "grand_livre") else "Octobre"  # fichier du mois passé
            contenu = ("4001", "Tremblay Immobilier inc.") if (code, doc) == ("4002", "sommaire") else (code, nom)
            creer_pdf(dossier / f"{nom.rstrip('.')}_{code}_{TITRES[doc]} de {mois}.pdf", *contenu, doc)  # 4002 : mauvais contenu
    creer_pdf(dossier / "Rapport annuel brouillon.pdf", "1001", "Boréale Immeubles inc.", "bilan")
    creer_pdf(dossier / "Entite inconnue inc_9999_Bilan de Octobre.pdf", "9999", "Entité inconnue", "bilan")


if __name__ == "__main__":
    creer_excel()
    creer_exports()
    print(f"Données fictives créées dans {ICI}")
