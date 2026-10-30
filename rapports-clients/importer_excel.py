"""Convertit le fichier de suivi « Clients contrôle mensuel.xlsx » au format de l'outil.

Lit l'onglet d'un mois (par défaut le dernier) et produit un Excel à quatre
onglets : Clients, Entites, Gestionnaires et « À vérifier ». Les entités qui
ont exactement les mêmes courriels clients sont regroupées sous un même client,
donc un seul courriel. Tout ce qui demande un jugement humain est listé dans
« À vérifier » au lieu d'être deviné.

    python importer_excel.py --source "Clients contrôle mensuel.xlsx" --sortie clients.xlsx
"""

import argparse
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

import tomllib

from envoi_rapports import ICI, cle_code, normaliser, numero_mois

CONTENU = tomllib.loads((ICI / "config.toml").read_text(encoding="utf-8"))["contenu"]
# Mots d'une note qui cochent un élément à inclure.
INDICES_NOTES = {
    "transactions_par_compte": ["transaction par", "transactions par"],
    "factures": ["factur"],
    "liste_cheques": ["cheque"],
    "liste_depots": ["depot"],
    "comptes_a_recevoir": ["compte a recevoir", "comptes a recevoir"],
    "sommaire_transactions": ["sommaire"],
}


def elements(note: str) -> dict[str, str]:
    note = normaliser(note)
    return {item["cle"]: "Oui" if item.get("defaut") or any(m in note for m in INDICES_NOTES.get(item["cle"], []))
            else "Non" for item in CONTENU}

# Courriels connus des gestionnaires (prénom tel qu'écrit dans le fichier de suivi).
GESTIONNAIRES_CONNUS = {"maxime": "mpaquet@progim.com"}


def nettoyer_courriel(valeur) -> str:
    return re.sub(r"^mailto:", "", str(valeur or "").strip(), flags=re.I).strip().lower()


def frequence(note: str, envoye: str) -> str:
    texte = normaliser(f"{note} {envoye}")
    if "trimestri" in texte or "30sep" in texte.replace(" ", ""):
        return "trimestriel"
    if "bi annuel" in texte or "semestri" in texte or "30juin" in texte.replace(" ", ""):
        return "semestriel"
    return "mensuel"


def lire_suivi(chemin: Path, onglet: str | None):
    classeur = load_workbook(chemin, data_only=True)
    feuille = classeur[onglet] if onglet else classeur.worksheets[-1]
    lignes = list(feuille.iter_rows(values_only=True))
    debut = next(i for i, l in enumerate(lignes) if l and normaliser(l[0] or "") == "code")
    entetes = [normaliser(c or "") for c in lignes[debut]]
    colonne = {nom: i for i, nom in enumerate(entetes)}
    courriels = [i for i, nom in enumerate(entetes) if nom.startswith("courriel client")]
    return feuille.title, lignes[debut + 1:], colonne, courriels


def convertir(source: Path, onglet: str | None, sortie: Path) -> None:
    titre, lignes, col, cols_courriels = lire_suivi(source, onglet)
    a_verifier: list[tuple[str, str]] = []
    entites: list[dict] = []

    for ligne in lignes:
        valeur = lambda nom: "" if ligne[col[nom]] is None else str(ligne[col[nom]]).strip()
        code, adresse = valeur("code"), valeur("adresses")
        emails = list(dict.fromkeys(e for e in (nettoyer_courriel(ligne[i]) for i in cols_courriels) if e))
        if not code:
            # Ligne sans code : immeuble supplémentaire de l'entité précédente.
            if adresse and entites:
                entites[-1]["immeubles"].append(adresse)
                if emails:
                    a_verifier.append((entites[-1]["code_entite"],
                                       f"Courriel(s) inscrits sur la ligne de l'immeuble « {adresse} », non repris : {', '.join(emails)}"))
            continue

        entite = {
            "code_entite": cle_code(code), "immeubles": [adresse] if adresse else [], "emails": emails,
            "fin_exercice": numero_mois(re.sub(r"^\d+\s*", "", valeur("fin d'annee fiscale"))),
            "frequence": frequence(valeur("note"), valeur("envoye au client")),
            "gestionnaire": "" if normaliser(valeur("gestionnaire")) in {"", "n/a"} else valeur("gestionnaire").title(),
            "notes": valeur("note"), "actif": "Oui", "contenu": elements(valeur("note")),
        }
        if "." not in code:
            a_verifier.append((entite["code_entite"], f"Code « {code} » sans décimale, interprété comme {entite['code_entite']}. "
                                                    "Vérifier le # dans le nom du PDF (ex. _08501)."))
        if entite["fin_exercice"] is None:
            a_verifier.append((entite["code_entite"], "Fin d'année fiscale non reconnue, 12 (décembre) par défaut."))
            entite["fin_exercice"] = 12
        if not emails:
            entite["actif"] = "Non"
            a_verifier.append((entite["code_entite"], "Aucun courriel client : entité mise inactive."))
        if entite["frequence"] != "mensuel":
            a_verifier.append((entite["code_entite"], f"Envoi {entite['frequence']} déduit de la note. À confirmer."))
        if any(m in normaliser(entite["notes"]) for m in ("ne plus envoyer", "factur", "combiner", "liste des transaction", "envoyer le")):
            a_verifier.append((entite["code_entite"], f"Cases cochées d'après la note « {entite['notes']} » : à vérifier."))
        entites.append(entite)

    # Un client = un ensemble identique de courriels clients.
    clients: dict[tuple, dict] = {}
    for entite in entites:
        if not entite["emails"]:
            cle = ("sans-courriel", entite["code_entite"])
        else:
            cle = tuple(sorted(entite["emails"]))
        if cle not in clients:
            clients[cle] = {"code_client": f"C{entite['code_entite']}", "entites": [], "emails": entite["emails"]}
        clients[cle]["entites"].append(entite)
        entite["code_client"] = clients[cle]["code_client"]

    for client in clients.values():
        noms = [e["code_entite"] for e in client["entites"]]
        client["actif"] = "Oui" if any(e["actif"] == "Oui" for e in client["entites"]) else "Non"
        if len(noms) > 1:
            a_verifier.append((client["code_client"], f"Regroupe {len(noms)} entités dans un seul courriel "
                                                     f"(mêmes destinataires) : {', '.join(e['code_entite'] for e in client['entites'])}."))
        if any(e.endswith("@progim.com") for e in client["emails"]):
            a_verifier.append((client["code_client"], "Courriel @progim.com parmi les destinataires clients. Voulu?"))

    a_verifier.insert(0, ("Tous", "nom_entite et nom_client sont vides : lancer envoi_rapports.py avec --completer sur "
                                  "les PDF du mois pour les remplir à partir des noms imprimés par Hopem, puis réviser."))
    gestionnaires = sorted({e["gestionnaire"] for e in entites if e["gestionnaire"]})
    for nom in gestionnaires:
        if normaliser(nom) not in GESTIONNAIRES_CONNUS:
            a_verifier.append(("Gestionnaires", f"Courriel de {nom} à inscrire."))

    classeur = Workbook()
    onglets = [
        ("Clients", ["code_client", "nom_client", "nom_contact", "courriel_clients", "courriel_direction", "langue", "actif"],
         [[c["code_client"], "", "", "; ".join(c["emails"]), "", "FR", c["actif"]] for c in clients.values()]),
        ("Entites", ["code_entite", "nom_entite", "immeubles", "code_client", "fin_exercice", "frequence",
                     "gestionnaire", "courriel_gestionnaire", "actif", *[i["cle"] for i in CONTENU], "notes"],
         [[e["code_entite"], "", "; ".join(e["immeubles"]), e["code_client"], e["fin_exercice"], e["frequence"],
           e["gestionnaire"], "", e["actif"], *e["contenu"].values(), e["notes"]] for e in entites]),
        ("Gestionnaires", ["nom", "courriel"],
         [[nom, GESTIONNAIRES_CONNUS.get(normaliser(nom), "")] for nom in gestionnaires]),
        ("A verifier", ["concerne", "a_verifier"], [list(v) for v in a_verifier]),
    ]
    entete, fond = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F5C4D")
    for i, (nom, colonnes, rangees) in enumerate(onglets):
        feuille = classeur.active if i == 0 else classeur.create_sheet()
        feuille.title = nom
        feuille.append(colonnes)
        for rangee in rangees:
            feuille.append(rangee)
        for cellule in feuille[1]:
            cellule.font, cellule.fill = entete, fond
        for colonne in feuille.columns:
            largeur = max(len(str(c.value or "")) for c in colonne) + 3
            feuille.column_dimensions[colonne[0].column_letter].width = min(largeur, 60)
            for c in colonne[1:]:
                c.number_format = "@"  # garde « 54.10 » et les codes comme du texte
        if rangees:
            table = Table(displayName="T_" + re.sub(r"\W", "", nom), ref=feuille.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            feuille.add_table(table)
        feuille.freeze_panes = "A2"
        if nom == "Entites":
            # Cases Oui/Non : liste déroulante sur l'actif et les éléments à inclure.
            choix = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
            feuille.add_data_validation(choix)
            for i, entete_col in enumerate(colonnes):
                if entete_col == "actif" or entete_col in {item["cle"] for item in CONTENU}:
                    lettre = feuille.cell(1, i + 1).column_letter
                    choix.add(f"{lettre}2:{lettre}{max(feuille.max_row, 200)}")
    classeur.save(sortie)

    actives = sum(e["actif"] == "Oui" for e in entites)
    print(f"Onglet « {titre} » : {len(entites)} entités ({actives} actives), {len(clients)} clients, "
          f"{len(a_verifier)} points à vérifier. Fichier : {sortie}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Convertit le fichier de suivi mensuel au format de l'outil.")
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--onglet", help="Onglet du mois à lire (défaut : le dernier)")
    parser.add_argument("--sortie", required=True, type=Path)
    args = parser.parse_args()
    convertir(args.source, args.onglet, args.sortie)
