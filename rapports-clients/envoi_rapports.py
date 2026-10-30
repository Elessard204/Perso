"""Préparation de l'envoi mensuel des rapports clients.

Lit l'Excel des clients, associe à chaque entité le PDF « Rapports mensuels »
téléchargé de Hopem, vérifie son contenu (nom, date, période, rapports
présents) et prépare un courriel par client. Rien n'est jamais envoyé : le
mode « brouillons » crée des brouillons dans la boîte finance@ que l'équipe
révise et envoie elle-même.

Exemples :
    python envoi_rapports.py --mois 2026-09 --clients clients.xlsx
    python envoi_rapports.py --mois 2026-09 --clients clients.xlsx --exports "C:/.../09-septembre 2026"
    python envoi_rapports.py --mois 2026-09 --clients clients.xlsx --exports ... --mode brouillons
"""

import argparse
import base64
import calendar
import html
import os
import re
import sys
import tomllib
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ICI = Path(__file__).resolve().parent
MOIS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
           "août", "septembre", "octobre", "novembre", "décembre"]
MOIS_EN = ["January", "February", "March", "April", "May", "June", "July",
           "August", "September", "October", "November", "December"]
# Débuts de mots reconnus pour un mois écrit en lettres (« Sept », « Fév. », « oct »...).
PREFIXES_MOIS = [("jan", 1), ("fe", 2), ("mar", 3), ("av", 4), ("ap", 4), ("mai", 5), ("may", 5),
                 ("juin", 6), ("jun", 6), ("juil", 7), ("jul", 7), ("ao", 8), ("aug", 8),
                 ("sep", 9), ("oct", 10), ("nov", 11), ("de", 12)]
# Mois civils où l'envoi a lieu, selon la fréquence de l'entité.
FREQUENCES = {"mensuel": set(range(1, 13)), "trimestriel": {3, 6, 9, 12}, "semestriel": {6, 12}}
COURRIEL_VALIDE = re.compile(r"^[^@\s;,]+@[^@\s;,]+\.[a-z]{2,}$", re.I)


def normaliser(texte) -> str:
    texte = unicodedata.normalize("NFKD", str(texte)).encode("ascii", "ignore").decode()
    return re.sub(r"[\s_\-]+", " ", texte).strip().lower()


def numero_mois(texte) -> int | None:
    texte = normaliser(texte).rstrip(".")
    return next((n for prefixe, n in PREFIXES_MOIS if texte.startswith(prefixe)), None)


def cle_code(code) -> str:
    """Code d'entité normalisé : « 54.1 », « 54,1 », « 054.01 » et « 05401 » donnent tous « 54.1 »."""
    texte = str(code).strip().replace(",", ".")
    if re.fullmatch(r"\d{5}", texte):
        return f"{int(texte[:3])}.{int(texte[3:])}"
    principal, _, sous = texte.partition(".")
    return f"{int(principal)}.{int(sous or 1)}" if principal.isdigit() and (sous or "1").isdigit() else texte


def oui(valeur) -> bool:
    return normaliser(valeur or "") in {"oui", "o", "x", "yes", "y", "1", "true", "vrai"}


def adresses(valeur) -> list[str]:
    morceaux = re.split(r"[;,\n]", str(valeur or ""))
    return [re.sub(r"^mailto:", "", a.strip(), flags=re.I) for a in morceaux if a.strip()]


def de_mois(numero: int) -> str:
    nom = MOIS_FR[numero - 1]
    return f"d'{nom}" if nom[0] in "aeiou" else f"de {nom}"


def mois_texte(mois: str, langue: str) -> str:
    noms = MOIS_EN if langue == "en" else MOIS_FR
    return f"{noms[int(mois[5:]) - 1]} {mois[:4]}"


@dataclass
class Entite:
    code: str
    nom: str          # nom légal, tel qu'imprimé par Hopem en haut du PDF
    immeubles: str
    code_client: str
    fin_exercice: int
    frequence: str
    gestionnaire_nom: str
    gestionnaires: list[str]
    actif: bool
    notes: str
    fichier: Path | None = None
    titres: list[str] = field(default_factory=list)
    nom_pdf: str = ""
    problemes: list[str] = field(default_factory=list)
    avertissements: list[str] = field(default_factory=list)

    def periode(self, mois: str) -> int:
        """Numéro de période Hopem du mois civil, selon la fin de l'exercice."""
        return (int(mois[5:]) - self.fin_exercice - 1) % 12 + 1

    def a_envoyer(self, mois: str) -> bool:
        mois_envoi = FREQUENCES.get(self.frequence, {self.fin_exercice})
        return self.actif and int(mois[5:]) in mois_envoi

    @property
    def nom_affiche(self) -> str:
        return self.nom or self.nom_pdf or self.immeubles

    def nom_piece(self, mois: str, langue: str) -> str:
        nom = re.sub(r'[\\/:*?"<>|]+', "-", self.nom_affiche).strip()
        return f"{nom} - {'Reports' if langue == 'en' else 'Rapports'} {mois_texte(mois, langue)}.pdf"


@dataclass
class Client:
    code: str
    nom: str
    nom_contact: str
    a: list[str]
    direction: list[str]
    langue: str
    actif: bool
    entites: list[Entite] = field(default_factory=list)
    problemes: list[str] = field(default_factory=list)
    objet: str = ""
    corps: str = ""
    brouillon: str = ""

    def a_envoyer(self, mois: str) -> list[Entite]:
        return [e for e in self.entites if e.a_envoyer(mois)] if self.actif else []

    def nom_affiche(self, mois: str) -> str:
        return self.nom or " et ".join(dict.fromkeys(e.nom_affiche for e in self.a_envoyer(mois) or self.entites))

    def pret(self, mois: str) -> bool:
        a_envoyer = self.a_envoyer(mois)
        return bool(a_envoyer) and not self.problemes and all(not e.problemes for e in a_envoyer)


def lire_onglet(classeur, nom: str) -> list[dict]:
    if nom not in classeur.sheetnames:
        return []
    lignes = list(classeur[nom].iter_rows(values_only=True))
    entetes = [normaliser(c).replace(" ", "_") if c else "" for c in lignes[0]]
    return [dict(zip(entetes, ligne)) for ligne in lignes[1:] if any(v not in (None, "") for v in ligne)]


def texte(valeur) -> str:
    return "" if valeur is None else str(valeur).strip()


def charger_clients(chemin: Path, config: dict) -> tuple[dict[str, Client], list[str]]:
    classeur = load_workbook(chemin, read_only=True, data_only=True)
    alertes = []

    courriels_gestionnaires = {
        normaliser(ligne["nom"]): adresses(ligne.get("courriel"))
        for ligne in lire_onglet(classeur, "Gestionnaires") if ligne.get("nom")
    }

    clients = {}
    for ligne in lire_onglet(classeur, "Clients"):
        code = texte(ligne["code_client"])
        if code in clients:
            alertes.append(f"Code client {code} en double dans l'Excel.")
        clients[code] = Client(
            code=code,
            nom=texte(ligne["nom_client"]),
            nom_contact=texte(ligne.get("nom_contact")),
            a=adresses(ligne["courriel_clients"]),
            direction=adresses(ligne.get("courriel_direction")) or config["envoi"]["direction_defaut"],
            langue="en" if normaliser(ligne.get("langue") or "fr") in {"en", "anglais"} else "fr",
            actif=oui(ligne.get("actif", "oui")),
        )

    codes_vus = set()
    for ligne in lire_onglet(classeur, "Entites"):
        code = cle_code(ligne["code_entite"])
        client = clients.get(texte(ligne["code_client"]))
        if client is None:
            alertes.append(f"Entité {code} rattachée au client « {texte(ligne['code_client'])} », absent de l'onglet Clients.")
            continue
        if code in codes_vus:
            alertes.append(f"Code d'entité {code} en double dans l'Excel.")
        codes_vus.add(code)
        try:
            fin = int(ligne.get("fin_exercice") or 12)
            assert 1 <= fin <= 12
        except (ValueError, AssertionError):
            alertes.append(f"Entité {code} : fin_exercice doit être un mois de 1 à 12.")
            continue
        frequence = normaliser(ligne.get("frequence") or "mensuel")
        if frequence not in FREQUENCES and frequence != "annuel":
            alertes.append(f"Entité {code} : fréquence « {frequence} » inconnue, traitée comme mensuelle.")
            frequence = "mensuel"
        gestionnaire = texte(ligne.get("gestionnaire"))
        gestionnaires = adresses(ligne.get("courriel_gestionnaire")) or courriels_gestionnaires.get(normaliser(gestionnaire), [])
        entite = Entite(
            code=code,
            nom=texte(ligne.get("nom_entite")),
            immeubles=texte(ligne.get("immeubles")),
            code_client=client.code,
            fin_exercice=fin,
            frequence=frequence,
            gestionnaire_nom=gestionnaire,
            gestionnaires=gestionnaires,
            actif=oui(ligne.get("actif", "oui")),
            notes=texte(ligne.get("notes")),
        )
        if gestionnaire and not gestionnaires:
            entite.avertissements.append(f"Gestionnaire « {gestionnaire} » sans courriel dans l'onglet Gestionnaires : pas en cc.")
        client.entites.append(entite)

    return clients, alertes


def associer_fichiers(dossier: Path, clients: dict[str, Client], mois: str, config: dict) -> list[str]:
    """Associe chaque PDF du dossier à son entité. Retourne les fichiers non classés."""
    motif = re.compile(config["fichiers"]["motif"], re.I)
    entites = {e.code: e for c in clients.values() for e in c.entites}
    non_classes = []

    for fichier in sorted(dossier.glob("*.pdf")):
        trouve = motif.match(fichier.name)
        if not trouve:
            non_classes.append(f"{fichier.name} : nom de fichier non reconnu.")
            continue
        entite = entites.get(cle_code(trouve["entite"]))
        if entite is None:
            non_classes.append(f"{fichier.name} : entité {cle_code(trouve['entite'])} introuvable dans l'Excel.")
            continue
        if not entite.a_envoyer(mois):
            raison = "inactive" if not entite.actif else f"envoi {entite.frequence}, pas prévu ce mois-ci"
            non_classes.append(f"{fichier.name} : entité {entite.code} {raison}. Fichier ignoré.")
            continue
        mois_fichier = numero_mois(trouve["mois"])
        annee_fichier = trouve.groupdict().get("annee")
        if mois_fichier != int(mois[5:]) or (annee_fichier and annee_fichier != mois[:4]):
            entite.problemes.append(f"{fichier.name} n'est pas le rapport {de_mois(int(mois[5:]))} {mois[:4]}.")
            continue
        if entite.fichier:
            entite.problemes.append(f"Deux fichiers pour la même entité : {entite.fichier.name} et {fichier.name}.")
            continue
        entite.fichier = fichier
    return non_classes


def lire_pdf(entite: Entite) -> str:
    """Relève le nom imprimé et les titres des rapports de chaque page. Retourne tout le texte."""
    tout = []
    for page in PdfReader(entite.fichier).pages:
        lignes = [l.strip() for l in (page.extract_text() or "").splitlines() if l.strip()]
        tout.extend(lignes)
        if lignes and re.match(r"^Page\s+\d+", lignes[0]):
            entite.nom_pdf = entite.nom_pdf or re.sub(r"^Page\s+\d+\s*", "", lignes[0])
            if len(lignes) > 2 and lignes[2] not in entite.titres:
                entite.titres.append(lignes[2])
    return "\n".join(tout)


def verifier_contenu(entite: Entite, mois: str, config: dict) -> None:
    controles = config["controles"]
    try:
        contenu = lire_pdf(entite)
    except Exception as erreur:
        entite.problemes.append(f"PDF illisible ({erreur}).")
        return

    titres = {normaliser(t) for t in entite.titres}
    for titre in controles["rapports_obligatoires"]:
        if normaliser(titre) not in titres:
            entite.problemes.append(f"{titre} absent du PDF.")

    annee, numero = int(mois[:4]), int(mois[5:])
    fin_mois = f"{annee}/{numero:02d}/{calendar.monthrange(annee, numero)[1]:02d}"
    date = re.search(r"\bAu\s*:?\s*(\d{4}/\d{2}/\d{2})", contenu)
    if controles.get("verifier_date") and date and date.group(1) != fin_mois:
        entite.problemes.append(f"Le bilan est au {date.group(1)}, attendu au {fin_mois}.")

    periode = re.search(r"p[ée]riode\s+(\d+)", contenu, re.I)
    if controles.get("verifier_periode") and periode and int(periode.group(1)) != entite.periode(mois):
        entite.problemes.append(
            f"L'état des résultats est pour la période {periode.group(1)}; avec une fin d'exercice "
            f"en {MOIS_FR[entite.fin_exercice - 1]}, on attend la période {entite.periode(mois)}. "
            "Vérifier la fin d'exercice dans l'Excel ou la période extraite.")

    # Protège contre un fichier mal numéroté : le nom imprimé par Hopem doit être celui attendu.
    if entite.nom and normaliser(entite.nom) != normaliser(entite.nom_pdf):
        entite.problemes.append(f"Le PDF est au nom de « {entite.nom_pdf} », attendu « {entite.nom} ». Mauvais fichier?")
    elif not entite.nom:
        entite.avertissements.append(f"Nom dans le PDF : « {entite.nom_pdf} ». À confirmer dans la colonne nom_entite.")


def valider(clients: dict[str, Client], mois: str, config: dict) -> None:
    for client in clients.values():
        a_envoyer = client.a_envoyer(mois)
        if not a_envoyer:
            continue
        if not client.a:
            client.problemes.append("Aucun courriel client.")
        for adresse in client.a + client.direction + [a for e in a_envoyer for a in e.gestionnaires]:
            if not COURRIEL_VALIDE.match(adresse):
                client.problemes.append(f"Adresse courriel invalide : {adresse}")
        for entite in a_envoyer:
            if entite.fichier:
                verifier_contenu(entite, mois, config)
            elif not entite.problemes:
                entite.problemes.append("Aucun PDF dans le dossier pour cette entité.")


def phrase_gestionnaires(noms: list[str], langue: str) -> str:
    if not noms:
        return "us" if langue == "en" else "nous"
    et = " and " if langue == "en" else " et "
    liste = ", ".join(noms[:-1]) + et + noms[-1] if len(noms) > 1 else noms[0]
    if langue == "en":
        return f"your property manager{'s' if len(noms) > 1 else ''}, {liste}, copied on this email"
    return f"{'vos gestionnaires' if len(noms) > 1 else 'votre gestionnaire'}, {liste}, en copie de ce courriel"


def ligne_periode(entite: Entite, mois: str, langue: str) -> str:
    annee, numero = int(mois[:4]), int(mois[5:])
    annee_fin = annee if numero <= entite.fin_exercice else annee + 1
    jour = calendar.monthrange(annee_fin, entite.fin_exercice)[1]
    if langue == "en":
        return f"period {entite.periode(mois)} of the fiscal year ending {MOIS_EN[entite.fin_exercice - 1]} {jour}, {annee_fin}"
    return f"période {entite.periode(mois)} de l'exercice se terminant le {jour} {MOIS_FR[entite.fin_exercice - 1]} {annee_fin}"


def rediger(client: Client, mois: str, config: dict) -> None:
    langue, entites = client.langue, client.a_envoyer(mois)
    gestionnaires = list(dict.fromkeys(e.gestionnaire_nom for e in entites if e.gestionnaires))
    titres = list(dict.fromkeys(t for e in entites for t in e.titres))
    valeurs = {
        "nom_client": client.nom_affiche(mois),
        "nom_contact": f" {html.escape(client.nom_contact)}" if client.nom_contact else "",
        "mois_texte": mois_texte(mois, langue),
        "gestionnaires": html.escape(phrase_gestionnaires(gestionnaires, langue)),
        "signature": html.escape(config["envoi"]["signature"]),
        "liste_entites": "<ul>" + "".join(
            f"<li>{html.escape(e.nom_affiche)}{f' ({html.escape(e.immeubles)})' if e.immeubles and e.nom_affiche != e.immeubles else ''}"
            f" <span style=\"color:#5d6b66\">· {ligne_periode(e, mois, langue)}</span></li>"
            for e in entites) + "</ul>",
        "liste_documents": html.escape(", ".join(t[0].lower() + t[1:] for t in titres)),
    }
    client.objet = config["envoi"][f"objet_{langue}"].format(**valeurs)
    modele = (ICI / "modeles" / f"courriel_{langue}.html").read_text(encoding="utf-8")
    client.corps = modele.format(**valeurs)


def destinataires_cc(client: Client, mois: str) -> list[str]:
    cc = [a for e in client.a_envoyer(mois) for a in e.gestionnaires] + client.direction
    deja = {a.lower() for a in client.a}
    resultat = []
    for adresse in cc:
        if adresse.lower() not in deja:
            deja.add(adresse.lower())
            resultat.append(adresse)
    return resultat


def ecrire_apercu(client: Client, mois: str, sortie: Path, config: dict) -> None:
    """Écrit un .eml qu'on peut ouvrir dans Outlook pour voir le courriel exact."""
    message = EmailMessage()
    message["From"] = config["envoi"]["expediteur"]
    message["To"] = ", ".join(client.a)
    message["Cc"] = ", ".join(destinataires_cc(client, mois))
    message["Subject"] = client.objet
    message["X-Unsent"] = "1"
    message.set_content("Ce courriel est en HTML.")
    message.add_alternative(client.corps, subtype="html")
    for entite in client.a_envoyer(mois):
        message.add_attachment(entite.fichier.read_bytes(), maintype="application", subtype="pdf",
                               filename=entite.nom_piece(mois, client.langue))
    dossier = sortie / "apercus"
    dossier.mkdir(parents=True, exist_ok=True)
    chemin = dossier / f"{client.code}.eml"
    chemin.write_bytes(bytes(message))
    client.brouillon = str(chemin.relative_to(sortie))


class Graph:
    """Création de brouillons dans la boîte de l'expéditeur via Microsoft Graph."""

    URL = "https://graph.microsoft.com/v1.0"
    LIMITE_PIECE_SIMPLE = 3 * 1024 * 1024

    def __init__(self, config: dict):
        import msal
        import requests

        self.requests = requests
        secret = os.environ.get("GRAPH_CLIENT_SECRET")
        if not (config["graph"]["tenant_id"] and config["graph"]["client_id"] and secret):
            sys.exit("Mode brouillons : tenant_id et client_id doivent être remplis dans config.toml "
                     "et la variable GRAPH_CLIENT_SECRET définie.")
        app = msal.ConfidentialClientApplication(
            config["graph"]["client_id"], client_credential=secret,
            authority=f"https://login.microsoftonline.com/{config['graph']['tenant_id']}")
        jeton = app.acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])
        if "access_token" not in jeton:
            sys.exit(f"Connexion à Microsoft 365 refusée : {jeton.get('error_description')}")
        self.entetes = {"Authorization": f"Bearer {jeton['access_token']}"}
        self.boite = f"{self.URL}/users/{config['envoi']['expediteur']}"

    def appel(self, methode: str, url: str, **kwargs) -> dict:
        reponse = self.requests.request(methode, url, headers=self.entetes, timeout=60, **kwargs)
        reponse.raise_for_status()
        return reponse.json() if reponse.content else {}

    def creer_brouillon(self, client: Client, mois: str) -> str:
        message = self.appel("POST", f"{self.boite}/messages", json={
            "subject": client.objet,
            "body": {"contentType": "HTML", "content": client.corps},
            "toRecipients": [{"emailAddress": {"address": a}} for a in client.a],
            "ccRecipients": [{"emailAddress": {"address": a}} for a in destinataires_cc(client, mois)],
        })
        for entite in client.a_envoyer(mois):
            self.joindre(message["id"], entite.fichier, entite.nom_piece(mois, client.langue))
        return message["webLink"]

    def joindre(self, id_message: str, fichier: Path, nom: str) -> None:
        contenu = fichier.read_bytes()
        url = f"{self.boite}/messages/{id_message}/attachments"
        if len(contenu) < self.LIMITE_PIECE_SIMPLE:
            self.appel("POST", url, json={
                "@odata.type": "#microsoft.graph.fileAttachment",
                "name": nom, "contentType": "application/pdf",
                "contentBytes": base64.b64encode(contenu).decode()})
            return
        # Les pièces de 3 Mo et plus passent par une session de téléversement.
        session = self.appel("POST", f"{url}/createUploadSession", json={"AttachmentItem": {
            "attachmentType": "file", "name": nom, "size": len(contenu)}})
        morceau = 4 * 1024 * 1024
        for debut in range(0, len(contenu), morceau):
            fin = min(debut + morceau, len(contenu))
            self.requests.put(session["uploadUrl"], data=contenu[debut:fin], timeout=120, headers={
                "Content-Length": str(fin - debut),
                "Content-Range": f"bytes {debut}-{fin - 1}/{len(contenu)}"}).raise_for_status()


def ecrire_rapport(clients: dict[str, Client], alertes: list[str], mois: str, mode: str,
                   sortie: Path, config: dict) -> Path:
    concernes = [c for c in clients.values() if c.a_envoyer(mois)]
    prets = [c for c in concernes if c.pret(mois)]
    bloques = [c for c in concernes if not c.pret(mois)]
    pas_ce_mois = [e for c in clients.values() if c.actif for e in c.entites if e.actif and not e.a_envoyer(mois)]
    e = html.escape

    def ligne_entite(x: Entite) -> str:
        details = []
        if x.frequence != "mensuel":
            details.append(x.frequence)
        if x.titres:
            details.append(" · ".join(x.titres))
        return (f'<li><span class="code">{e(x.code)}</span> {e(x.nom_affiche)} <span class="periode">P{x.periode(mois)}</span>'
                + (f' <span class="discret">· {e(" · ".join(details))}</span>' if details else "")
                + (f'<div class="discret">Immeubles : {e(x.immeubles)}</div>' if x.immeubles else "")
                + (f'<div class="discret">Note : {e(x.notes)}</div>' if x.notes else "")
                + "".join(f'<div class="probleme">{e(p)}</div>' for p in x.problemes)
                + "".join(f'<div class="avertissement">{e(a)}</div>' for a in x.avertissements)
                + "</li>")

    def ligne_client(c: Client) -> str:
        pret = c.pret(mois)
        etat = '<span class="pastille ok">Prêt</span>' if pret else '<span class="pastille bloque">Bloqué</span>'
        entites = "".join(ligne_entite(x) for x in c.a_envoyer(mois))
        problemes = "".join(f'<div class="probleme">{e(p)}</div>' for p in c.problemes)
        lien = f'<div class="discret">Aperçu : <span class="code">{e(c.brouillon)}</span></div>' if c.brouillon else ""
        gestionnaires = ", ".join(dict.fromkeys(x.gestionnaire_nom for x in c.a_envoyer(mois) if x.gestionnaire_nom))
        return (f'<tr class="{"" if pret else "a-voir"}"><td>{etat}</td>'
                f'<td><strong>{e(c.nom_affiche(mois))}</strong><div class="discret code">{e(c.code)}</div></td>'
                f'<td><div>{e(", ".join(c.a) or "—")}</div><div class="discret">gestionnaire : {e(gestionnaires or "aucun")}</div></td>'
                f'<td><ul class="entites">{entites}</ul>{problemes}{lien}</td></tr>')

    # Liste d'extraction : quelle période sortir de Hopem pour chaque entité.
    par_periode: dict[int, list[Entite]] = {}
    for c in concernes:
        for x in c.a_envoyer(mois):
            par_periode.setdefault(x.periode(mois), []).append(x)
    extraction = "".join(
        f'<div class="groupe"><div class="groupe-tete"><span class="periode grande">Période {p}</span>'
        f'<span class="discret">{len(xs)} entité{"s" if len(xs) > 1 else ""} · fin d\'exercice en '
        f'{", ".join(sorted({MOIS_FR[x.fin_exercice - 1] for x in xs}))}</span></div>'
        f'<div class="discret">{" · ".join(f"<span class=code>{e(x.code)}</span> {e(x.nom_affiche)}" for x in xs)}</div></div>'
        for p, xs in sorted(par_periode.items()))

    apercu = ""
    if prets and mode != "liste":
        c = prets[0]
        apercu = f"""
<section>
  <h2>Aperçu d'un courriel préparé</h2>
  <div class="courriel">
    <dl>
      <dt>De</dt><dd>{e(config["envoi"]["expediteur"])}</dd>
      <dt>À</dt><dd>{e(", ".join(c.a))}</dd>
      <dt>Cc</dt><dd>{e(", ".join(destinataires_cc(c, mois)))}</dd>
      <dt>Objet</dt><dd><strong>{e(c.objet)}</strong></dd>
      <dt>Pièces</dt><dd>{"".join(f'<span class="piece">{e(x.nom_piece(mois, c.langue))}</span>' for x in c.a_envoyer(mois))}</dd>
    </dl>
    <div class="corps">{c.corps}</div>
  </div>
</section>"""

    hors_excel = "".join(f"<li>{e(a)}</li>" for a in alertes)
    pas_ce_mois_html = " · ".join(f'<span class="code">{e(x.code)}</span> {e(x.nom_affiche)} ({e(x.frequence)})' for x in pas_ce_mois)
    libelle_mode = {"brouillons": "Brouillons créés dans Outlook",
                    "simulation": "Simulation : aucun brouillon créé",
                    "liste": "Liste d'extraction : à faire avant de télécharger les rapports de Hopem"}[mode]
    page = f"""<!doctype html>
<html lang="fr"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<title>Envoi des rapports {e(mois_texte(mois, "fr"))}</title>
<meta name="claude:viewport-fit-sides" content="cover">
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Public+Sans:wght@400;600;700&family=IBM+Plex+Mono:wght@400;500&display=swap">
<style>
/* Mise en page : une colonne de contrôle, sommaire d'abord, détail par client ensuite. */
:root {{
  --fond: #f6f7f5; --surface: #ffffff; --texte: #1d2522; --discret: #5d6b66;
  --ligne: #dde3df; --accent: #1f5c4d; --ok: #24734f; --ok-fond: #e3f2ea;
  --bloque: #a33a1f; --bloque-fond: #fbe9e3; --attention: #8a5a00;
  --police: "Public Sans", system-ui, sans-serif; --mono: "IBM Plex Mono", ui-monospace, monospace;
}}
@media (prefers-color-scheme: dark) {{ :root:not([data-theme="light"]) {{
  --fond: #121715; --surface: #1a211e; --texte: #e4ebe7; --discret: #9aa8a2;
  --ligne: #2c3632; --accent: #7cc4ad; --ok: #6fd0a0; --ok-fond: #183326;
  --bloque: #f0957a; --bloque-fond: #3a1f17; --attention: #e5b85c; color-scheme: dark; }} }}
:root[data-theme="dark"] {{
  --fond: #121715; --surface: #1a211e; --texte: #e4ebe7; --discret: #9aa8a2;
  --ligne: #2c3632; --accent: #7cc4ad; --ok: #6fd0a0; --ok-fond: #183326;
  --bloque: #f0957a; --bloque-fond: #3a1f17; --attention: #e5b85c; color-scheme: dark; }}
[hidden] {{ display: none !important; }}
body {{ margin: 0; background: var(--fond); color: var(--texte); font: 15px/1.5 var(--police); }}
main {{ max-width: 1100px; margin: 0 auto; padding-block: 32px 64px;
  padding-left: max(16px, env(safe-area-inset-left, 0px)); padding-right: max(16px, env(safe-area-inset-right, 0px));
  display: grid; gap: 32px; }}
header {{ display: grid; gap: 6px; }}
.surtitre {{ font: 500 12px var(--mono); letter-spacing: .08em; text-transform: uppercase; color: var(--accent); }}
h1 {{ margin: 0; font-size: clamp(24px, 4vw, 34px); line-height: 1.15; text-wrap: balance; }}
h2 {{ margin: 0 0 12px; font-size: 18px; }}
.discret {{ color: var(--discret); font-size: 13px; }}
.code {{ font-family: var(--mono); font-size: 12.5px; }}
.compteurs {{ display: flex; flex-wrap: wrap; gap: 12px; }}
.compteur {{ flex: 1 1 160px; background: var(--surface); border: 1px solid var(--ligne); border-radius: 6px; padding: 14px 16px; }}
.compteur b {{ display: block; font-size: 30px; font-variant-numeric: tabular-nums; line-height: 1.1; }}
.compteur.ok b {{ color: var(--ok); }} .compteur.bloque b {{ color: var(--bloque); }}
.tableau {{ overflow-x: auto; background: var(--surface); border: 1px solid var(--ligne); border-radius: 6px; }}
table {{ width: 100%; border-collapse: collapse; min-width: 760px; }}
th {{ text-align: left; font: 500 11.5px var(--mono); letter-spacing: .06em; text-transform: uppercase; color: var(--discret);
  padding: 10px 14px; border-bottom: 1px solid var(--ligne); }}
td {{ padding: 12px 14px; border-bottom: 1px solid var(--ligne); vertical-align: top; }}
td:last-child {{ width: 50%; }}
tr:last-child td {{ border-bottom: 0; }}
tr.a-voir td:first-child {{ box-shadow: inset 3px 0 0 var(--bloque); }}
.pastille {{ display: inline-block; font-size: 12px; font-weight: 600; padding: 2px 9px; border-radius: 99px; white-space: nowrap; }}
.pastille.ok {{ color: var(--ok); background: var(--ok-fond); }}
.pastille.bloque {{ color: var(--bloque); background: var(--bloque-fond); }}
.entites {{ list-style: none; margin: 0; padding: 0; display: grid; gap: 8px; }}
.probleme {{ color: var(--bloque); font-size: 13px; }}
.probleme::before {{ content: "→ "; }}
.avertissement {{ color: var(--attention); font-size: 13px; }}
.avertissement::before {{ content: "· "; }}
.periode {{ font: 500 11.5px var(--mono); color: var(--accent); border: 1px solid var(--ligne); border-radius: 4px; padding: 0 5px; }}
.periode.grande {{ font-size: 13px; padding: 1px 8px; }}
.extraction {{ display: grid; gap: 1px; background: var(--ligne); border: 1px solid var(--ligne); border-radius: 6px; overflow: hidden; }}
.groupe {{ background: var(--surface); padding: 12px 16px; display: grid; gap: 4px; }}
.groupe-tete {{ display: flex; flex-wrap: wrap; align-items: baseline; gap: 4px 12px; }}
.alertes {{ margin: 0; padding-left: 20px; color: var(--bloque); }}
.courriel {{ background: var(--surface); border: 1px solid var(--ligne); border-radius: 6px; overflow: hidden; }}
.courriel dl {{ margin: 0; padding: 14px 18px; display: grid; grid-template-columns: max-content 1fr; gap: 6px 16px;
  border-bottom: 1px solid var(--ligne); font-size: 14px; }}
.courriel dt {{ color: var(--discret); }} .courriel dd {{ margin: 0; min-width: 0; overflow-wrap: anywhere; }}
.piece {{ display: inline-block; font: 12.5px var(--mono); border: 1px solid var(--ligne); border-radius: 4px; padding: 1px 8px; margin: 0 6px 4px 0; }}
.corps {{ padding: 8px 18px 14px; max-width: 65ch; }}
</style></head>
<body><main>
<header>
  <div class="surtitre">Finance · rapports clients</div>
  <h1>Envoi des rapports {e(de_mois(int(mois[5:])))} {mois[:4]}</h1>
  <div class="discret">{libelle_mode} · préparé le {datetime.now():%Y-%m-%d à %H:%M}</div>
</header>
<div class="compteurs"{" hidden" if mode == "liste" else ""}>
  <div class="compteur ok"><b>{len(prets)}</b>courriels prêts</div>
  <div class="compteur bloque"><b>{len(bloques)}</b>courriels bloqués</div>
  <div class="compteur"><b>{sum(len(c.a_envoyer(mois)) for c in prets)}</b>rapports d'entités joints</div>
  <div class="compteur"><b>{len(pas_ce_mois)}</b>entités sans envoi ce mois-ci</div>
</div>
<section>
  <h2>Périodes à extraire dans Hopem</h2>
  <div class="extraction">{extraction}</div>
</section>
{f'<section><h2>Fichiers ou lignes à régler</h2><ul class="alertes">{hors_excel}</ul></section>' if alertes else ""}
<section{" hidden" if mode == "liste" else ""}>
  <h2>Détail par client</h2>
  <div class="tableau"><table>
    <thead><tr><th>État</th><th>Client</th><th>Destinataires</th><th>Entités et vérifications</th></tr></thead>
    <tbody>{"".join(ligne_client(c) for c in bloques + prets)}</tbody>
  </table></div>
</section>
{f'<section><h2>Pas d&#39;envoi ce mois-ci</h2><div class="discret">{pas_ce_mois_html}</div></section>' if pas_ce_mois else ""}
{apercu}
</main></body></html>"""
    chemin = sortie / "rapport.html"
    chemin.write_text(page, encoding="utf-8")
    return chemin


def completer_excel(chemin: Path, clients: dict[str, Client], mois: str) -> Path:
    """Copie l'Excel en remplissant les noms vides avec ceux trouvés dans les PDF, pour révision."""
    classeur = load_workbook(chemin)
    trouves = {e.code: e for c in clients.values() for e in c.entites if e.nom_pdf}
    remplis = 0
    for nom_onglet, cle, colonne_nom in (("Entites", "code_entite", "nom_entite"), ("Clients", "code_client", "nom_client")):
        feuille = classeur[nom_onglet]
        entetes = [normaliser(c.value or "").replace(" ", "_") for c in feuille[1]]
        if colonne_nom not in entetes:
            continue
        i_cle, i_nom = entetes.index(cle), entetes.index(colonne_nom)
        for rangee in feuille.iter_rows(min_row=2):
            if rangee[i_nom].value or rangee[i_cle].value is None:
                continue
            if nom_onglet == "Entites":
                entite = trouves.get(cle_code(rangee[i_cle].value))
                nom = entite.nom_pdf if entite else ""
            else:
                client = clients.get(texte(rangee[i_cle].value))
                nom = " et ".join(dict.fromkeys(e.nom_pdf for e in client.entites if e.nom_pdf)) if client else ""
            if nom:
                rangee[i_nom].value = nom
                remplis += 1
    copie = chemin.with_name(f"{chemin.stem}_complete.xlsx")
    classeur.save(copie)
    print(f"{remplis} noms ajoutés à partir des PDF, à réviser : {copie}")
    return copie


def main() -> None:
    parser = argparse.ArgumentParser(description="Prépare l'envoi mensuel des rapports clients.")
    parser.add_argument("--mois", required=True, help="Mois des rapports, format AAAA-MM")
    parser.add_argument("--clients", required=True, type=Path, help="Excel des clients")
    parser.add_argument("--exports", type=Path,
                        help="Dossier des rapports téléchargés de Hopem. Sans ce dossier, l'outil produit "
                             "seulement la liste des périodes à extraire.")
    parser.add_argument("--sortie", type=Path, help="Dossier de sortie (défaut : ./sortie/AAAA-MM)")
    parser.add_argument("--mode", choices=["simulation", "brouillons"], default="simulation")
    parser.add_argument("--config", type=Path, default=ICI / "config.toml")
    parser.add_argument("--completer", action="store_true",
                        help="Écrit une copie de l'Excel où les noms vides sont remplis à partir des PDF.")
    args = parser.parse_args()

    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", args.mois):
        sys.exit("--mois doit être au format AAAA-MM, par exemple 2026-09.")
    config = tomllib.loads(args.config.read_text(encoding="utf-8"))
    sortie = args.sortie or ICI / "sortie" / args.mois
    sortie.mkdir(parents=True, exist_ok=True)

    clients, alertes = charger_clients(args.clients, config)
    if args.exports is None:
        rapport = ecrire_rapport(clients, alertes, args.mois, "liste", sortie, config)
        print(f"Liste des périodes à extraire : {rapport}")
        return
    alertes += associer_fichiers(args.exports, clients, args.mois, config)
    valider(clients, args.mois, config)

    if args.completer:
        completer_excel(args.clients, clients, args.mois)

    graph = Graph(config) if args.mode == "brouillons" else None
    for client in clients.values():
        if not client.pret(args.mois):
            continue
        rediger(client, args.mois, config)
        ecrire_apercu(client, args.mois, sortie, config)
        if graph:
            client.brouillon = graph.creer_brouillon(client, args.mois)

    rapport = ecrire_rapport(clients, alertes, args.mois, args.mode, sortie, config)
    concernes = [c for c in clients.values() if c.a_envoyer(args.mois)]
    print(f"{sum(c.pret(args.mois) for c in concernes)} courriels prêts, "
          f"{sum(not c.pret(args.mois) for c in concernes)} bloqués, {len(alertes)} alertes. Rapport : {rapport}")


if __name__ == "__main__":
    main()
